import tkinter as tk
from tkinter import messagebox, ttk
import openpyxl
import os

FILE = "student_results.xlsx"


# Create Excel file
if not os.path.exists(FILE):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Name", "Roll No.", "Class", "Sub1", "Sub2", "Sub3",
               "Sub4", "Sub5", "Total", "Percentage", "Result"])
    wb.save(FILE)


# Add Student
def add_student():
    win = tk.Toplevel(root)
    win.title("Add Student")
    win.geometry("400x600")
    win.config(bg="lightgrey")

    labels = ["Name", "Roll No.", "Class", "Subject 1", "Subject 2",
              "Subject 3", "Subject 4", "Subject 5"]

    entries = []

    for label in labels:
        tk.Label(
            win,
            text=label,
            font=("Arial", 12),
            bg="lightgrey"
        ).pack(pady=5)

        e = tk.Entry(
            win,
            font=("Arial", 12)
        )
        e.pack()

        entries.append(e)

    def save():
        name = entries[0].get()
        roll = entries[1].get()
        student_class = entries[2].get()

        marks = []

        for e in entries[3:]:
            marks.append(float(e.get()))

        total = sum(marks)
        percentage = total / 5

        if all(mark >= 40 for mark in marks):
            result = "Pass"
        else:
            result = "Fail"

        wb = openpyxl.load_workbook(FILE)
        ws = wb.active

        ws.append([
            name,
            roll,
            student_class,
            *marks,
            total,
            percentage,
            result
        ])

        wb.save(FILE)

        messagebox.showinfo(
            "Success",
            "Student saved successfully!"
        )

        win.destroy()

    tk.Button(
        win,
        text="💾 Save",
        font=("Arial", 12),
        command=save
    ).pack(pady=20)


# Get Result
def get_result():
    win = tk.Toplevel(root)
    win.title("Get Result")
    win.geometry("500x400")
    win.config(bg="lightgrey")

    tk.Label(
        win,
        text="Enter Roll No.",
        font=("Arial", 14),
        bg="lightgrey"
    ).pack(pady=15)

    entry = tk.Entry(
        win,
        font=("Arial", 12)
    )
    entry.pack()

    result_label = tk.Label(
        win,
        text="",
        font=("Arial", 13),
        bg="lightgrey"
    )
    result_label.pack(pady=30)

    def search():
        roll = entry.get()

        wb = openpyxl.load_workbook(FILE)
        ws = wb.active

        for row in ws.iter_rows(min_row=2, values_only=True):

            if str(row[1]) == roll:

                result_label.config(
                    text=f"Name: {row[0]}\n"
                         f"Roll No.: {row[1]}\n"
                         f"Class: {row[2]}\n"
                         f"Total Marks: {row[8]}\n"
                         f"Percentage: {row[9]:.1f}%\n"
                         f"Result: {row[10]}"
                )

                wb.close()
                return

        result_label.config(
            text="❌ Student record not found."
        )

        wb.close()

    tk.Button(
        win,
        text="🔍 Get Result",
        font=("Arial", 12),
        command=search
    ).pack(pady=10)


# Show All Results
def show_all():
    win = tk.Toplevel(root)
    win.title("All Results")
    win.geometry("800x400")
    win.config(bg="lightgrey")

    tk.Label(
        win,
        text="All Student Results",
        font=("Arial", 18, "bold"),
        bg="lightgrey"
    ).pack(pady=15)

    columns = (
        "Name",
        "Roll No.",
        "Class",
        "Total",
        "Percentage",
        "Result"
    )

    tree = ttk.Treeview(
        win,
        columns=columns,
        show="headings"
    )

    for col in columns:
        tree.heading(col, text=col)
        tree.column(col, width=120)

    tree.pack(
        fill="both",
        expand=True,
        padx=10,
        pady=10
    )

    wb = openpyxl.load_workbook(FILE)
    ws = wb.active

    for row in ws.iter_rows(min_row=2, values_only=True):

        tree.insert(
            "",
            "end",
            values=(
                row[0],
                row[1],
                row[2],
                row[8],
                f"{row[9]:.1f}%",
                row[10]
            )
        )

    wb.close()


# Main Window
root = tk.Tk()
root.title("Student Result Management System")
root.geometry("500x450")
root.config(bg="lightgrey")


tk.Label(
    root,
    text="STUDENT RESULT MANAGEMENT",
    font=("Arial", 20, "bold"),
    bg="lightgrey"
).pack(pady=40)


tk.Button(
    root,
    text="1. Add Student",
    width=25,
    font=("Arial", 13),
    command=add_student
).pack(pady=10)


tk.Button(
    root,
    text="2. Get Result",
    width=25,
    font=("Arial", 13),
    command=get_result
).pack(pady=10)


tk.Button(
    root,
    text="3. Show All Results",
    width=25,
    font=("Arial", 13),
    command=show_all
).pack(pady=10)


tk.Button(
    root,
    text="4. Exit",
    width=25,
    font=("Arial", 13),
    command=root.destroy
).pack(pady=10)


root.mainloop()